import json, datetime
from collections import Counter, defaultdict
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.comments import Comment

I='/tmp/claude-0/inv/'
fichas=json.load(open(I+'fichas-5550.json'))
ten=json.load(open(I+'tenedores-5550.json'))
hue=json.load(open(I+'huellas.json'))
rol=json.load(open(I+'roles-5550.json'))
pools=json.load(open(I+'pools-5550.json'))
emp=json.load(open(I+'emparejado-8532.json'))
ordx=json.load(open(I+'ordenex-mercados.json'))
verif=set(open(I+'verificados.txt').read().split())
cread=json.load(open('/home/user/express-js-on-vercel/infra/migracion-cadena/creadores-gestores.json'))
esc=json.load(open(I+'5550/fase1-conteos.json'))
CORTE=esc['corte']; FECHA_CORTE='2026-09-22'
DESPLEGADORA='0xa07a2ba9735568b989378bbb3b1439ee9c5e5714'
ND='NO DISPONIBLE'
f_by={f['direccion']:f for f in fichas}
dup=Counter(f['sha256'] for f in fichas)
simcount=Counter((f['simbolo'] or '') for f in fichas if f['simbolo'])
sym=lambda a: (f_by.get(a,{}).get('simbolo') or f_by.get(a,{}).get('nombre') or a[:10]) if a else ''

def unidades(base, dec):
    if base is None or dec is None: return None
    return int(base)/(10**int(dec))

FUENTE=Font(name='Arial',size=10); NEG=Font(name='Arial',size=10,bold=True); TIT=Font(name='Arial',size=14,bold=True)
CAB=Font(name='Arial',size=10,bold=True,color='FFFFFF'); FCAB=PatternFill('solid',fgColor='0E1A2B')
FND=PatternFill('solid',fgColor='FBECD1'); FEXTRA=PatternFill('solid',fgColor='EEF2F7'); FALERTA=PatternFill('solid',fgColor='F6DCD8')
fina=Side(style='thin',color='D8DCE3')
wb=Workbook()

# ───────────── Léeme ─────────────
ws=wb.active; ws.title='Léeme'
filas=[
 ('Inventario de contratos · Bloque 1 de la solicitud SFSP v0.2',None),
 ('',None),
 ('Corte de lectura',f'Red 5550, bloque {CORTE:,} · {FECHA_CORTE}'.replace(',','.')),
 ('Cómo se obtuvo','Lectura directa del nodo público rpc.ordenglobal-rpc.com (Besu 26.7.1), sólo consultas: eth_call, eth_getCode, eth_getStorageAt, recibos y trazas. No se firmó ni se movió nada.'),
 ('Contratos en la 5550','172. Coincide con los 172 contratos con código que registró el constructor del génesis. Se recorrió la cadena entera: 241.307 bloques, 95 transacciones, 0 despliegues de contratos (directos o internos).'),
 ('Por qué todos son del bloque 0','La 5550 actual arrancó el 25-ago-2026 (bloque 1: 23:40:15 UTC) con un génesis que copió código y almacenamiento de la cadena anterior. Ningún contrato se desplegó después. Fuente: infra/migracion-cadena/reinicio-5550/ETAPA-3-EL-CORTE.md; el hash del bloque 0 del nodo (0x56b3cf56…) coincide con el del acta.'),
 ('Tenedores','Se preguntó el saldo de 465 direcciones conocidas (código, datos de la migración, eventos). Si la suma de saldos iguala el supply total, la lista está COMPLETA y la concentración es exacta. Para AUKA, ONDK, HARV y USDK se leyeron además, por su huella en el almacenamiento, los saldos de direcciones que no están en el repositorio.'),
 ('NO DISPONIBLE','Casilla en amarillo: el dato no se puede obtener de la cadena ni del repositorio. No se estimó nada.'),
 ('Columnas extra','Casillas en gris-azul a la derecha de las pedidas: hechos que ayudan a clasificar (huella del bytecode, duplicados, funciones privilegiadas presentes, cobertura de tenedores).'),
 ('La cadena 8532','No respondió el 22-sep (23.23.205.33:80, conexión rechazada). La hoja 1.2 sale del volcado completo del estado al corte (genesis-id/datos/cadena-8532.json, origen s3://ordenglobal-cadena-8532-respaldo/2026-08-10) emparejado con la 5550 dirección por dirección y, donde faltaba la dirección, por su huella keccak.'),
 ('Alertas','Ver la hoja «Alertas»: hallazgos que afectan el diseño de la migración.'),
]
for r,(a,b) in enumerate(filas,1):
    ws.cell(r,1,a).font = TIT if r==1 else NEG
    if b: c=ws.cell(r,2,b); c.font=FUENTE; c.alignment=Alignment(wrap_text=True,vertical='top')
ws.column_dimensions['A'].width=26; ws.column_dimensions['B'].width=120
ws.cell(9,2).fill=FEXTRA; ws.cell(8,2).fill=FND

# ───────────── 1.1 ─────────────
ws=wb.create_sheet('1.1 Contratos 5550')
PEDIDAS=['Dirección del contrato','Nombre','Símbolo','Estándar','Decimales','Supply total','Supply circulante','Bloque de despliegue','Fecha de despliegue',
 'Cuenta desplegadora','Roles privilegiados','Número de tenedores','Concentración (top 10)','Última transferencia','Código verificado','Contratos relacionados','Responsable interno','Origen']
EXTRA=['Suma top 10 (unidades)','Cobertura de tenedores','Supply sin ubicar (unidades)','En poder de la cuenta desplegadora (unidades)','Funciones privilegiadas en el bytecode','Contratos con el mismo bytecode','SHA-256 del bytecode','Huella registrada en ORDENSCAN','Indicios (hechos, no clasificación)','Alerta']
hdr=PEDIDAS+EXTRA
for j,h in enumerate(hdr,1):
    c=ws.cell(1,j,h); c.font=CAB; c.fill=FCAB; c.alignment=Alignment(wrap_text=True,vertical='center')
ws.freeze_panes='B2'
# qué contratos aparecen en qué pools
en_pools=defaultdict(list)
for p in pools:
    for t in (p['token0'],p['token1']):
        if t: en_pools[t].append(p['direccion'])
alertas=[]
orden=sorted(fichas,key=lambda f:(not f['estandar'].startswith('ERC-20'), f['estandar'], f['simbolo'] or 'zzz', f['direccion']))
fila=2
for f in orden:
    a=f['direccion']; t=ten.get(a); h=hue.get(a); dec=f['decimales']
    caps=[k for k in ('mint','burn','burnFrom','pause','unpause','transferOwnership','renounceOwnership','upgradeTo','blacklist') if f['capacidades'].get(k)]
    roles=[]
    if f['owner']: roles.append(f"owner(): {f['owner']}")
    if f['impl']: roles.append(f"proxy · implementación: {f['impl']}")
    if f['adminProxy']: roles.append(f"proxy · administrador: {f['adminProxy']}")
    roles_txt='; '.join(roles) if roles else ('Sin dueño legible por owner(); ' + (f"funciones presentes: {', '.join(caps)} · quién puede llamarlas: {ND} (sin fuente)" if caps else 'sin funciones privilegiadas detectadas'))
    # tenedores
    ntxt=None; top10=None; cob=None; sinubicar=None; despl=None; alerta=''
    if t:
        vals=[int(v) for _,v in t['tenedores']]
        extra=[int(v) for v in (h['valoresDesconocidos'] if h else [])]
        todos=sorted(vals+extra,reverse=True)
        n=len(todos)
        if f['estandar']=='ERC-721':
            ntxt=f"{len(vals)} (con saldo entre direcciones conocidas)"
        elif t['completo'] or (h and h['cierra']):
            ntxt=n
            if extra: ntxt=f"{n} (exacto; {len(extra)} sin dirección conocida, leídos por huella)"
        else:
            ntxt=(f"{len(vals)} con dirección conocida + hasta {len(extra)} entradas sin dirección (pueden incluir autorizaciones)" if extra else f"al menos {len(vals)}")
        if dec is not None: top10=sum(todos[:10])/(10**int(dec))
        sup=int(t['supply']) if t['supply'] not in (None,'0') else None
        if sup:
            ident=sum(vals)+sum(extra)
            cob=ident/sup
            if ident!=sup and dec is not None: sinubicar=(sup-ident)/(10**int(dec))
        dv=dict(t['tenedores']).get(DESPLEGADORA)
        if dv and dec is not None: despl=int(dv)/(10**int(dec))
    if f['simbolo']=='WORIGEN':
        alerta='totalSupply()=0 pero balanceOf(cuenta desplegadora)=55.000.000; el contrato no tiene ORIGEN. Saldos sin respaldo.'
    if f['nombre']=='Wrapped Origen' and f['simbolo']=='WETH':
        alerta='Envoltorio vaciado en el génesis del 25-ago (acordado): 0 ORIGEN y 0 tenedores. Los pools que lo usan declaran liquidez con 0 WETH real.'
    if alerta: alertas.append((a, f['simbolo'] or f['nombre'] or f['estandar'], alerta))
    ult=rol['ultima'].get(a)
    ult_txt=datetime.datetime.utcfromtimestamp(ult).strftime('%Y-%m-%d %H:%M UTC') if ult else f'Ninguna desde el génesis (25-ago-2026). Historial en la 8532: {ND}'
    rel=[]
    if f['token0']: rel.append(f"token0 {sym(f['token0'])} {f['token0']}")
    if f['token1']: rel.append(f"token1 {sym(f['token1'])} {f['token1']}")
    if f['factory']: rel.append(f"factoría {f['factory']}")
    if f['impl']: rel.append(f"implementación {f['impl']}")
    if en_pools.get(a): rel.append(f"en {len(en_pools[a])} par(es)/pool(s) AMM")
    creador = cread.get(a)
    desp_txt = f"Ninguna en la 5550 (sembrado en el génesis). En la 8532: {creador[0]} (nonce {creador[1]}; infra/migracion-cadena/creadores-gestores.json)" if creador else f"Ninguna en la 5550 (sembrado en el génesis). En la 8532: {ND}"
    ind=[]
    if a in verif: ind.append('en el catálogo de ORDENSCAN')
    if f['simbolo'] and simcount[f['simbolo']]>1: ind.append(f"el símbolo {f['simbolo']} aparece en {simcount[f['simbolo']]} contratos")
    if dup[f['sha256']]>1: ind.append(f"bytecode idéntico a otros {dup[f['sha256']]-1}")
    if f['factoriaPares']=='0': ind.append('factoría sin ningún par creado')
    vals_row=[a, f['nombre'] or ND, f['simbolo'] or ND, f['estandar'], int(dec) if dec is not None else ND,
      unidades(t['supply'] if t else f['supply'], dec) if (t or f['supply']) and dec is not None else (ND if f['estandar'].startswith('ERC') else 'No aplica'),
      f'{ND}: falta la lista oficial de billeteras de tesorería y control interno (ver columnas extra)' if f['estandar'].startswith('ERC-20') else 'No aplica',
      0, '2026-08-25 (génesis de la 5550; bloque 1 a las 23:40:15 UTC). Despliegue original en la 8532: '+ND,
      desp_txt, roles_txt, ntxt if ntxt is not None else 'No aplica', None, ult_txt,
      'No. El repositorio no tiene la fuente Solidity; ver «Huella registrada en ORDENSCAN»', '; '.join(rel) or '—', ND, ND,
      top10, cob, sinubicar, despl, ', '.join(caps) or '—', dup[f['sha256']]-1, f['sha256'], 'Sí' if a in verif else 'No', '; '.join(ind) or '—', alerta or '—']
    for j,v in enumerate(vals_row,1):
        c=ws.cell(fila,j,v); c.font=FUENTE; c.alignment=Alignment(vertical='top',wrap_text=j in (11,14,16,18,27,28))
        if isinstance(v,str) and v.startswith(ND): c.fill=FND
        if j>len(PEDIDAS): c.fill=FEXTRA
        if j==len(hdr) and alerta: c.fill=FALERTA
    # concentración como fórmula: top10 / supply
    colS=get_column_letter(6); colT=get_column_letter(len(PEDIDAS)+1)
    if top10 is not None and isinstance(vals_row[5],(int,float)) and vals_row[5]>0:
        ws.cell(fila,13,f'={colT}{fila}/{colS}{fila}').number_format='0.00%'
        if (cob is not None and cob<1) or (h and h['valoresDesconocidos'] and not h['cierra']):
            ws.cell(fila,13).comment=Comment('Aproximada: faltan tenedores por ubicar y hay entradas sin dirección que podrían ser autorizaciones y no saldos (ver «Cobertura de tenedores» y la hoja Tenedores).','inventario')
    else:
        ws.cell(fila,13,'No aplica' if not f['estandar'].startswith('ERC-20') else ND)
    ws.cell(fila,13).font=FUENTE
    for j in (6,19,21,22): ws.cell(fila,j).number_format='#,##0.####'
    ws.cell(fila,20).number_format='0.0000%'
    fila+=1
anch=[44,24,12,22,9,20,30,9,30,40,50,22,14,30,30,50,14,12,18,14,18,20,34,10,66,12,50,50]
for j,w in enumerate(anch,1): ws.column_dimensions[get_column_letter(j)].width=w
ws.auto_filter.ref=f'A1:{get_column_letter(len(hdr))}{fila-1}'
tot=fila
ws.cell(tot+1,1,'Total de contratos').font=NEG; ws.cell(tot+1,2,f'=COUNTA(A2:A{fila-1})').font=NEG
for k,(e,n) in enumerate(sorted(Counter(f['estandar'] for f in fichas).items(), key=lambda x:-x[1])):
    ws.cell(tot+2+k,1,e).font=FUENTE; ws.cell(tot+2+k,2,f'=COUNTIF(D2:D{fila-1},A{tot+2+k})').font=FUENTE

# ───────────── Tenedores ─────────────
ws=wb.create_sheet('Tenedores')
for j,h in enumerate(['Token','Contrato','Tenedor','Saldo (unidades)','% del supply','¿Contrato?','Nota'],1):
    c=ws.cell(1,j,h); c.font=CAB; c.fill=FCAB
r=2
for f in orden:
    a=f['direccion']; t=ten.get(a)
    if not t or not f['estandar'].startswith('ERC-20') or f['decimales'] is None: continue
    dec=int(f['decimales']); sup=int(t['supply']) if t['supply'] not in (None,'0') else None
    filas_t=[(x,int(v),'') for x,v in t['tenedores']]
    h=hue.get(a)
    if h: filas_t+= [(f'(sin dirección conocida #{i+1})',int(v),'leído por su huella en el almacenamiento; podría ser una autorización y no un saldo') for i,v in enumerate(h['valoresDesconocidos'])]
    filas_t.sort(key=lambda x:-x[1])
    ini=r
    for x,v,nota in filas_t:
        ws.cell(r,1,f['simbolo']); ws.cell(r,2,a); ws.cell(r,3,x); ws.cell(r,4,v/10**dec).number_format='#,##0.####'
        ws.cell(r,5,f'=D{r}/INDEX(\'1.1 Contratos 5550\'!F:F,MATCH(B{r},\'1.1 Contratos 5550\'!A:A,0))' if sup else 'No aplica').number_format='0.0000%'
        ws.cell(r,6,'Sí' if x in f_by else ('Desplegadora' if x==DESPLEGADORA else 'No'))
        ws.cell(r,7,nota)
        for j in range(1,8): ws.cell(r,j).font=FUENTE
        r+=1
for j,w in enumerate([12,44,44,22,12,12,60],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='A2'; ws.auto_filter.ref=f'A1:G{r-1}'

# ───────────── 1.2 ─────────────
ws=wb.create_sheet('1.2 Contratos 8532')
h2=['Dirección del contrato (8532)','Nombre','Símbolo','Estándar','Decimales','Supply total','Supply circulante','Bloque de despliegue','Fecha de despliegue','Cuenta desplegadora','Roles privilegiados','Número de tenedores','Concentración','Última transferencia','Código verificado','Contratos relacionados','Responsable interno','Origen','Estado respecto de la migración','Dirección en la 5550','Saldo nativo al corte (ORIGEN)','Cómo se identificó la dirección']
for j,hh in enumerate(h2,1):
    c=ws.cell(1,j,hh); c.font=CAB; c.fill=FCAB; c.alignment=Alignment(wrap_text=True)
r=2
for e in sorted(emp,key=lambda x:(x['en5550'] is None, x['en5550'] or '')):
    a5=e['en5550']; f=f_by.get(a5,{})
    estado='Migrado (misma dirección, código y almacenamiento copiados al génesis de la 5550)' if a5 else 'Abandonado a propósito: contrato de staking de Polygon Edge; en la 5550 los validadores los gobierna QBFT (GENESIS-5550-LISTO.md)'
    ver='Ver hoja 1.1 (mismo código y almacenamiento; puede diferir por transacciones posteriores al corte)'
    vals=[e['direccion8532'] or a5, f.get('nombre') or ND, f.get('simbolo') or ND, f.get('estandar') or ND, int(f['decimales']) if f.get('decimales') else ND,
          ver if a5 else ND, ver if a5 else ND, ND, ND, (f"{cread[a5][0]} (creadores-gestores.json)" if a5 in cread else ND), ver if a5 else ND, ver if a5 else ND, ver if a5 else ND,
          ND, 'No (sin fuente en el repositorio)', ver if a5 else '—', ND, ND, estado, a5 or '—', int(e['saldoNativo'])/1e18,
          'por su huella keccak (el volcado no traía la dirección)' if e['recuperadaPorHuella'] else ('volcado del estado' if e['direccion8532'] else '—')]
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=FUENTE
        if isinstance(v,str) and v.startswith(ND): c.fill=FND
    ws.cell(r,21).number_format='#,##0.######'
    r+=1
for j,w in enumerate([44,24,12,22,9,30,30,10,10,40,30,30,30,12,26,30,12,10,60,44,18,34],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='B2'
ws.cell(r+1,1,'Contratos en la 8532').font=NEG; ws.cell(r+1,2,f'=COUNTA(A2:A{r-1})')
ws.cell(r+2,1,'Migrados').font=NEG; ws.cell(r+2,2,f'=COUNTIF(S2:S{r-1},"Migrado*")')
ws.cell(r+3,1,'Abandonados').font=NEG; ws.cell(r+3,2,f'=COUNTIF(S2:S{r-1},"Abandonado*")')

# ───────────── 1.3 ─────────────
ws=wb.create_sheet('1.3 Mercados')
h3=['Par negociado','Dónde','Dirección del contrato','Estado operativo','Fecha de creación','Volumen histórico acumulado','Fecha de la última operación','Liquidez hoy','Nota']
for j,hh in enumerate(h3,1):
    c=ws.cell(1,j,hh); c.font=CAB; c.fill=FCAB; c.alignment=Alignment(wrap_text=True)
r=2
for m in ordx:
    vals=[m['mercado'],'Ordenex (libro de órdenes fuera de la cadena, MongoDB)','No aplica: no es un contrato',
          'Abierto' + (' · sólo hay orden de venta' if m.get('mejorVenta') and not m.get('mejorCompra') else ' · sin órdenes'),
          ND, '0 (API pública: /mercados/'+m['mercado']+'/tratos devuelve cero operaciones)', 'Ninguna operación registrada', '—',
          f"Referencia: {m['referencia'].get('rotulo','')} · fuente {m['referencia'].get('fuente','')}" if m.get('referencia') else '—']
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=FUENTE; c.alignment=Alignment(wrap_text=True,vertical='top')
        if isinstance(v,str) and v.startswith(ND): c.fill=FND
    r+=1
for p in sorted(pools,key=lambda x:(x['tipo'],x['par'])):
    liq = (p['reserva0'] not in (None,'0')) or (p['liquidezV3'] not in (None,'0'))
    vals=[p['par'], p['tipo']+' en la cadena 5550', p['direccion'], 'Con liquidez' if liq else 'Sin liquidez',
          '2026-08-25 en la 5550 (génesis). Creación original en la 8532: '+ND, f'{ND}: los eventos de intercambio de la 8532 no pasaron al génesis', 'Ninguna en la 5550', 
          (f"reservas {p['reserva0']} / {p['reserva1']} (unidades base)" if p['reserva0'] is not None else f"liquidez V3 {p['liquidezV3']}"),
          'Saldo real de WETH del pool: 0 (ver Alertas)' if 'WETH' in p['par'] and liq else '—']
    for j,v in enumerate(vals,1):
        c=ws.cell(r,j,v); c.font=FUENTE; c.alignment=Alignment(wrap_text=True,vertical='top')
        if isinstance(v,str) and ND in v: c.fill=FND
    r+=1
for j,w in enumerate([24,34,44,26,34,40,24,40,40],1): ws.column_dimensions[get_column_letter(j)].width=w
ws.freeze_panes='A2'

# ───────────── Alertas ─────────────
ws=wb.create_sheet('Alertas')
AL=[
 ('Supply de AUKA y ONDK no ubicado del todo','AUKA: faltan 9.823,01 unidades. ONDK: faltan 792,5 unidades. Ni entre las 465 direcciones conocidas ni entre las entradas de almacenamiento de la foto de la 8532. Probablemente son saldos creados en la 5550 intermedia (15 al 25 de agosto), cuyo historial se perdió con el reinicio. No está probado. Para ubicarlos hace falta la base de usuarios de Veta Wallet o el volcado de esa cadena intermedia.'),
 ('Tres contratos WORIGEN incoherentes','totalSupply() = 0, pero balanceOf(cuenta desplegadora 0xa07a…5714) = 55.000.000 en los tres, y ninguno tiene ORIGEN. Si alguien los canjeara, no hay respaldo. Hay que decidir su tratamiento antes de migrar.'),
 ('Pools contra WETH con liquidez que no existe','El envoltorio 0xccbe…9f75 tenía 16.387,76 ORIGEN en la 8532. El génesis del 25-ago lo vació junto con sus saldos, como estaba acordado (reinicio-5550/ETAPAS-1-Y-2.md). Pero 12 pools V3 contra WETH siguen declarando liquidez interna con 0 WETH real: sólo conservan su lado en token (por ejemplo, unos 11,39 AUKA entre tres pools). Esa liquidez no se puede operar. Los tokens que quedan dentro pertenecen, según el acta, a las dos cuentas operadoras que crearon las posiciones (0x0186450c…, 0x3063a26b…). Hay que decidir si se retiran antes de migrar. El único pool con los dos lados reales es TKNA/TKNB.'),
 ('Ningún código fuente','El repositorio no tiene la fuente Solidity de ninguno de los 172 contratos. «Código verificado» es No en todos. ORDENSCAN registra sólo la huella del bytecode de 15.'),
 ('Despliegues repetidos','24 factorías con el mismo bytecode (la mayoría sin pares), varias copias de routers, envoltorios y contratos mínimos. Es un indicio de despliegues de prueba, no una clasificación: el campo «Origen» queda NO DISPONIBLE hasta que alguien responsable lo confirme.'),
 ('La 8532 no responde','El 21-ago respondía en 23.23.205.33:80 (LA-8532-SIGUE-VIVA.md); el 22-sep rechaza la conexión. Si hace falta algún dato de la 8532 que no esté en el volcado, se necesita el respaldo en S3.'),
]
for k,(a,b) in enumerate(AL,1):
    ws.cell(k,1,a).font=NEG; c=ws.cell(k,2,b); c.font=FUENTE; c.alignment=Alignment(wrap_text=True,vertical='top')
k=len(AL)+2
ws.cell(k,1,'Por contrato').font=NEG
for a,s,al in alertas:
    k+=1; ws.cell(k,1,s).font=FUENTE; ws.cell(k,2,f'{a} · {al}').font=FUENTE
ws.column_dimensions['A'].width=38; ws.column_dimensions['B'].width=130

wb.save('/home/user/express-js-on-vercel/sfsp/respuestas/BLOQUE-1-INVENTARIO.xlsx')
print('ok', fila-2, 'contratos 5550;', len(emp),'contratos 8532;', len(ordx)+len(pools),'mercados')
